"""업로드된 파일에서 텍스트를 뽑아내는 유틸리티 모음.

지원 형식: PDF(.pdf), Word(.docx), PowerPoint(.pptx), 이미지(.png/.jpg/.jpeg/.gif/.webp)
※ 예전 바이너리 형식인 .doc, .ppt는 지원하지 않습니다 (OOXML 형식만 지원).
"""

import os

from pypdf import PdfReader
from docx import Document
from pptx import Presentation

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}


def get_file_kind(filename: str) -> str:
    """파일명(확장자)을 보고 'pdf' / 'docx' / 'pptx' / 'image' / 'unsupported' 중 하나를 반환합니다."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext == ".pptx":
        return "pptx"
    if ext in IMAGE_EXTENSIONS:
        return "image"
    return "unsupported"


def extract_text_from_pdf(path: str) -> str:
    reader = PdfReader(path)
    parts = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(parts).strip()


def extract_text_from_docx(path: str) -> str:
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    # 표 안의 텍스트도 함께 뽑아준다
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(cell.text.strip())
    return "\n".join(lines).strip()


def extract_text_from_pptx(path: str) -> str:
    prs = Presentation(path)
    lines = []
    for i, slide in enumerate(prs.slides, start=1):
        lines.append(f"[슬라이드 {i}]")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text.strip():
                        lines.append(text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                lines.append(f"(발표자 노트: {notes.strip()})")
    return "\n".join(lines).strip()
